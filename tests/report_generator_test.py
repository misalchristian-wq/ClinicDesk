"""Smoke-test the official workbook export against its supplied template."""
import json
import subprocess
import sys
import tempfile
from pathlib import Path

import openpyxl


ROOT = Path(__file__).resolve().parents[1]
TEMPLATE = ROOT / "report_generator" / "SCHOOL_YEAR_REPORT.xlsx"
GENERATOR = ROOT / "report_generator" / "generate_deped_report.py"


def main():
    sample = {
        "school": {"name":"Synthetic School","id":"123456"},
        "students": [{"school_year": "2025-2026", "grade_level": "7", "sex": "Female", "bmi_category": "Wasted"}],
        "immunization": [{"grade_level": "7", "sex": "Female", "vaccine": "Tetanus Diphtheria", "immunized": 1}],
        "deworming": [{"lrn": "123", "school_year": "2025-2026", "grade_level": "7", "sex": "Female", "dewormed_sbfp": 1, "dewormed_other": 0, "wifa": 0, "wifa_date": None}],
        "arh": [{"grade_level": "7", "delivery_mode": "In School", "total": 1}],
        "peer_educators": 2,
        "tobacco": [{"level_group": "jhs", "brought": 1, "referred": 1}],
        "lhas": [{"grade_level": "7", "screening_type": "Vision Screening", "masterlisted": 1, "screened": 1, "findings": 0, "referred_school": 0, "referred_lgu": 0, "referred_private": 0, "referred_others": 0}],
        "saved_reports": {
            "box2_3": {"hasSchoolClinic": "Yes", "visitedBySDO": "Yes", "sdoVisits": 2, "clinicEquipment": {"First Aid Kit": "Functional"}, "waterSources": ["Water Well"], "waterForDrinking": "No"},
            "box1": {"referralConcerns": ["Medical, dental and nutritional"]},
            "box4": {"hasGuidanceOffice": "Yes", "counselingJHS": {"male": 2, "female": 3}, "vulnerableJHS": {"muslim": 1}, "hasMentalHealthTraining": "Yes", "mentalHealthTraining": {"bullying": 4}, "mentalHealthCases":{"attemptInside":{"jhsLearners":1}}},
            "box5_6": {"hasSupportCenter": "Yes", "iecMaterials": ["No Smoking Signages"], "storesSelling": ["Tobacco products"], "tobaccoIntervention":{"jhs":{"users":2,"bti":1}}},
            "box8_9": {"drugEducation":"Yes","drugComponents":["Curriculum integration"],"drugLifeSkills":{"g7":4},"hasCanteen": "Yes", "canteenManager": "Others", "canteenManagerOther": "PTA", "sanitaryPermit": "Yes", "healthCertificates": "No", "hasKitchen": "Yes", "feedingFundSources": ["SBFP"], "agriResources": ["Fish Pond"]},
            "box10_11": {"swmImplementation": ["Composting"], "stakeholders": ["Parents"], "sanitaryPadLocations": ["Others"], "sanitaryPadOther": "Office"},
        },
    }
    for grade in range(7, 13):
        for date in ("2025-08-14", "2026-02-14"):
            sample["deworming"].append({"lrn": str(1000 + grade), "school_year": "2025-2026", "grade_level": str(grade), "sex": "Female", "wifa": 1, "wifa_date": date})
    # A second dose in the same period and a date outside the school year must not add to the count.
    sample["deworming"].extend([
        {"lrn": "1007", "school_year": "2025-2026", "grade_level": "7", "sex": "Female", "wifa": 1, "wifa_date": "2025-09-01"},
        {"lrn": "9999", "school_year": "2025-2026", "grade_level": "7", "sex": "Female", "wifa": 1, "wifa_date": "2026-08-01"},
    ])
    with tempfile.TemporaryDirectory() as directory:
        data_path = Path(directory) / "report.json"
        output_path = Path(directory) / "report"
        data_path.write_text(json.dumps(sample), encoding="utf-8")
        result = subprocess.run([sys.executable, str(GENERATOR), str(TEMPLATE), str(output_path), "2025-2026", str(data_path)], capture_output=True, text=True, check=True)
        assert json.loads(result.stdout)["success"]
        with output_path.open("rb") as generated_file:
            book = openpyxl.load_workbook(generated_file)
        assert "2025-2026" in book["IX. Box 1"]["A2"].value
        assert book["Part IX. School Health"]["D5"].value == "Synthetic School - 123456"
        assert book["IX. Table 1.A & 1.B"]["L37"].value == 1
        assert book["IX. Table 1.C & 1.D"]["Q21"].value == 1
        wifa_sheet = book["IX. Table 1.C & 1.D"]
        assert "July to September 2025" in wifa_sheet["C35"].value
        assert "January to March 2026" in wifa_sheet["C37"].value
        for col in ("Q", "V", "AA", "AF", "AP", "AU"):
            assert wifa_sheet[f"{col}35"].value == 1
            assert wifa_sheet[f"{col}37"].value == 1
        assert book["IX. Boxes 2 & 3; Table 2"]["AA6"].value == "/"
        assert book["IX. Boxes 2 & 3; Table 2"]["Z39"].value is None
        assert book["IX. Boxes 2 & 3; Table 2"]["U18"].value == "/"
        assert book["IX. Boxes 2 & 3; Table 2"]["Z18"].value is None
        assert 'clinic_sdo_personnel_yes="/"' in book["IX. Boxes 2 & 3; Table 2"]["Z10"].value
        assert book["IX. Boxes 2 & 3; Table 2"]["N10"].value == 2
        assert book["IX. Boxes 2 & 3; Table 2"]["X60"].value == 1
        assert book["IX. Box 1"]["U8"].value == "/"
        assert book["IX. Box 1"]["U9"].value is None
        assert book["IX. Box 4"]["M14"].value == 2
        assert book["IX. Boxes 5 & 6"]["V13"].value == "/"
        assert book["IX. Boxes 5 & 6"]["Z49"].value == 2
        assert book["IX. Boxes 7 to 9"]["Z6"].value == "/"
        assert book["IX. Boxes 7 to 9"]["D17"].value == 4
        assert book["IX. Boxes 7 to 9"]["M30"].value == "PTA"
        assert book["IX. Boxes 7 to 9"]["T52"].value == "/"
        assert book["IX. Boxes 7 to 9"]["W32"].value == "/"
        assert book["IX. Boxes 7 to 9"]["W34"].value is None
        assert book["IX. Boxes 10 & 11"]["AB8"].value == "/"
        assert book["IX. Boxes 10 & 11"]["AB9"].value is None
        assert book["IX. Boxes 10 & 11"]["L37"].value == "Office"
        assert not any(isinstance(cell.value, bool) for sheet in book if sheet.title.startswith("IX.") for row in sheet for cell in row)
        assert book.sheetnames == openpyxl.load_workbook(TEMPLATE, read_only=True).sheetnames
        sample["saved_reports"] = {}
        data_path.write_text(json.dumps(sample), encoding="utf-8")
        result = subprocess.run([sys.executable, str(GENERATOR), str(TEMPLATE), str(output_path), "2025-2026", str(data_path)], capture_output=True, text=True, check=True)
        assert json.loads(result.stdout)["success"]
        with output_path.open("rb") as generated_file:
            partial = openpyxl.load_workbook(generated_file)
        assert partial["IX. Table 1.A & 1.B"]["L37"].value == 1
        assert partial["IX. Box 4"]["M14"].value is None
    print("Report generator template check passed")


if __name__ == "__main__":
    main()
